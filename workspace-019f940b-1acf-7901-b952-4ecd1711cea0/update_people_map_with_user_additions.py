from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from copy import copy

src='South_Africa_SAP_CAR_POSDTA_people_market_map.xlsx'
out='South_Africa_SAP_CAR_POSDTA_people_market_map_v2.xlsx'
wb=load_workbook(src)

header_fill=PatternFill('solid', fgColor='1F4E78')
header_font=Font(color='FFFFFF', bold=True)
green=PatternFill('solid', fgColor='E2F0D9')
yellow=PatternFill('solid', fgColor='FFF2CC')
orange=PatternFill('solid', fgColor='FCE4D6')
grey=PatternFill('solid', fgColor='E7E6E6')
blue=PatternFill('solid', fgColor='D9EAF7')
thin=Side(style='thin', color='D9D9D9')
border=Border(left=thin,right=thin,top=thin,bottom=thin)

def append_row(ws, row, fill):
    ws.append(row)
    r=ws.max_row
    for c in range(1, ws.max_column+1):
        cell=ws.cell(r,c)
        cell.alignment=Alignment(wrap_text=True, vertical='top')
        cell.border=border
        cell.fill=fill

# Add high-confidence new rows to SA-linked map
ws=wb['SA-linked people map']
new_rows=[
    ['Tristan','Bell','SAP Functional Analyst','The Shoprite Group of Companies','Malmesbury / Brackenfell, Western Cape','https://www.linkedin.com/in/ttb16072000/','Yes','POSDM / sales audit adjacent; POSDTA not named','No direct evidence','No direct evidence','No direct evidence','Direct LinkedIn profile states he specializes in SAP CAR POSDM; configures and integrates systems for NOVA Point of Sale rollout across multiple countries; enhances sales audit processes and maintains reports from a large enterprise database.','A — direct LinkedIn profile evidence','Direct LinkedIn profile surfaced in public search','https://www.linkedin.com/in/ttb16072000/','Profile says POSDM rather than POSDTA. Because SAP CAR POSDM may be legacy wording, do not rewrite as POSDTA unless separately confirmed.'],
    ['Stasha','Battye','Vice President Group Digital Payments / payments and POS programme leader','Massmart / Block Markets Africa (current LinkedIn result varies by date)','City of Johannesburg, Gauteng','https://za.linkedin.com/in/stasha-battye-41bbaa9','Adjacent SAP/POS ecosystem','Yes — SAP POS / POS programme evidence, not POSDTA','No direct evidence','No direct evidence','No direct evidence','Direct LinkedIn profile describes 25+ years in retail POS/payments; UCS roles include service owner for SAP POS, project manager for SAP All-in-One and Active Retail integration services, and technical PM for Makro POS equipment/system integration rollout.','B — direct LinkedIn evidence, adjacent not CAR/POSDTA','Direct LinkedIn profile surfaced in public search','https://za.linkedin.com/in/stasha-battye-41bbaa9','Added as adjacent POS/SAP Retail project/program manager. No SAP CAR/POSDTA/DDF/UDF/OAA evidence found.']
]
for row in new_rows:
    append_row(ws,row,green if row[12].startswith('A') else yellow)

# Add user additions triage sheet
if 'User additions triage' in wb.sheetnames:
    del wb['User additions triage']
ws=wb.create_sheet('User additions triage')
headers=['Name','User-supplied LinkedIn URL','User-supplied role/employer','CAR/POSDTA/DDF/UDF/OAA evidence found in this sweep','Decision for CAR/POSDTA people map','Verification status / notes']
ws.append(headers)
for c in range(1,len(headers)+1):
    cell=ws.cell(1,c); cell.fill=header_fill; cell.font=header_font; cell.alignment=Alignment(wrap_text=True, vertical='top'); cell.border=border

def add(name,url,role,ev,decision,notes,fill):
    ws.append([name,url,role,ev,decision,notes])
    r=ws.max_row
    for c in range(1,7):
        cell=ws.cell(r,c); cell.alignment=Alignment(wrap_text=True, vertical='top'); cell.border=border; cell.fill=fill

# Verified/checked in this turn
add('Tristan Bell','https://www.linkedin.com/in/ttb16072000/','SAP Functional Analyst — Shoprite','Direct LinkedIn evidence: SAP CAR POSDM, NOVA POS rollout, sales audit processes.','ADDED to SA-linked people map as direct SAP CAR/POSDM specialist.','Verified from direct LinkedIn public result. POSDTA not named; kept as POSDM/CAR wording.',green)
add('Stasha Battye','https://za.linkedin.com/in/stasha-battye-41bbaa9','VP Group Digital Payments — Massmart / Block Markets Africa','Direct LinkedIn evidence: SAP POS, POS programme management, Makro POS replacement, SAP Retail integration. No CAR/POSDTA/DDF/UDF/OAA.','ADDED as adjacent POS/SAP Retail programme leader, not strict CAR/POSDTA.','Verified from direct LinkedIn public result. Useful for POS ecosystem, not CAR/POSDTA.',yellow)
add('Nolan Reddy','https://www.linkedin.com/in/nolanreddy/','User supplied: SAP Delivery Manager / Project Manager — Woolworths','No accepted evidence found. Searches surfaced other Nolan Reddy profiles unrelated to Woolworths/SAP CAR.','NOT ADDED.','Claim not verified; public results contradicted or did not match supplied URL/role.',orange)

# Names from user list not yet accepted; grouped out of scope/pending
names_general = [
'Arno Botha','Janel Rust','Jenny Bath','Sean Brown','Sango Ndiki','Patrick Crowe','Amanda Strydom','Ryogil Tafadzwa Mangarayi','Annelize De Beer','Nicky De Jager','Jerome Samson','Charl Gouws','Sibusiso Fakude','Paul Obino','Penuel Nkosi','Siddeeqah Stevens','Alan Joshua','Pieter Kellerman','Tebogo Mokele','Zanele Khumalo','Sarita Pillay','Johan Coetzee','Fatima Ismail','Thembeka Cele','Jacques Roux','Mbali Ndlovu','Luan van der Merwe','Sonwabo Ndiki','Kevin Naidu','Sipho Gumede','Lerato Molefe','Andre van Wyk','Preveshan Pillay','Thando Maseko','Liezl Meyer','Bongani Sibanda','Karen Smith','Neo Matlou','David Brown','Michael Petersen','Samantha Govender','Chris van den Berg','Nonhlanhla Dlamini','Riaan Swart','Jessica Lee','Mandla Zulu','Elize Coetzee','Brandon Naidoo','Sabelo Khumalo','Rajesh Kumar Roy','Souhail Samadi','Maxim Behm','Lukasz Kwiatkowski','Bosega Nchaupe','Duba Selepe','Brett Schreuder','Kumar Shaankar Anand','Thiren Ramkisson','Ram Dodda','Malene Marie Bendixen Jacobsen','Nokukhanya Mbali Dlamini','Nizamuddin Shaikh','Bongiwe Sibanda','Busisiwe Mnisi','Sagar Kumar Jha','Antima Sharma','Marilene Joubert','Bongani Hlungwane','Ntebo Mofokeng','Nqobile Mabaso','Ayanda Zebe','Themba Maseko','Siyabonga Ntombela','Jason Sewrajan','Ramadimetja Moshapalo','Sabelo Mnguni','Samkelo Nkambule','Derek Breetzke','Nthabiseng Mashiane','Sifiso Msabala','Sphamandla Ndlovu','Prelene Parkash','Letlotlo Baloyi','Siboniso Khumalo','Daniel Engelbrecht','Kobus Steenkamp','Sabelo Ntsaluba','Bennie Groenewald','Neil Slabbert','Rebecca Steyn','Sibalo Dlamini','Chris Moller','Daleen Dudley','Ezz El-Din Mamdouh','Gerrit Du Plessis','Trevor Gouvas','Edgar Mhlanga','Aveshnee Moopanar','Dziripi Godo','Mpume Zondi','Kabelo Lefifi','Moses Morapedi Maduna','Seelan Nair','Ruben Chetty','Elicus Marais','Dumisani Moyo','Melucy Mlotshwa','Sascha Perkuhn','Jinchandra Mule','Peter Siddle','Lerato Mokoena','Wayne Govender','Shobhit Kumar','Sachin Kapoor','Akhilesh Bhat','Soukaina Benayad','Avinash M','Sapna Naik','Morgan Naicker','Anusha Pillay','Nhlanhla Mahlangu','Tatenda Tarusarira','Jean Olivier','Jerry De Lange','Pedro Guerreiro','Simon Van Wyk','Lindiwe Mahlangu','Mamoupo Regina Masilo','Jabulani Zwane','Tshepo Mokoka','Sifiso Sdumo','Petr Dlouhy','Jimmy Van Geit','Skyos Van Niekerk','Leslie Nkosi'
]
for name in names_general:
    add(name,'Provided by user; not re-used as verified source','General SAP / SAP PM / Agile / adjacent role from user-supplied list','No accepted CAR/POSDTA/DDF/UDF/OAA evidence identified in this turn.','NOT ADDED to strict CAR/POSDTA people map yet.','These may be useful for a separate broad South Africa SAP project-manager/functional-consultant database, but they require individual source verification and most are outside CAR/POSDTA/DDF/UDF/OAA scope.',grey)

# Update summary
wsum=wb['Executive summary']
row=wsum.max_row+2
wsum.cell(row,1).value='User-supplied additions processed'
wsum.cell(row,1).font=Font(bold=True); wsum.cell(row,1).fill=blue; wsum.cell(row,1).border=border
wsum.cell(row,2).value='Added Tristan Bell as a direct SAP CAR/POSDM specialist and Stasha Battye as adjacent SAP POS/SAP Retail programme leader. Nolan Reddy/Woolworths SAP CAR claim was not verified. Remaining user-supplied names are triaged as general SAP/PM/adjacent and need separate verification before any inclusion in a broadened SAP PM market map.'
wsum.cell(row,2).alignment=Alignment(wrap_text=True, vertical='top'); wsum.cell(row,2).border=border

# Formatting
for wsname in wb.sheetnames:
    wsx=wb[wsname]
    for col in range(1, wsx.max_column+1):
        if wsx.column_dimensions[get_column_letter(col)].width is None or wsx.column_dimensions[get_column_letter(col)].width < 12:
            wsx.column_dimensions[get_column_letter(col)].width=22
    if wsx.max_row>1:
        wsx.freeze_panes='A2'
        wsx.auto_filter.ref=wsx.dimensions

wb.save(out)
print(out)
