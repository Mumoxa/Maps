from __future__ import annotations

import csv
import hashlib
import json
import re
import sys
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


PLACEHOLDERS = {
    "", "-", "n/a", "na", "none", "not found", "not available", "unknown",
    "#ref!", "#n/a", "false", "tbd",
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = re.sub(r"\s+", " ", str(value)).strip()
    return "" if text.casefold() in PLACEHOLDERS else text


def normalise(value: str) -> str:
    text = unicodedata.normalize("NFKD", clean(value)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", text.casefold()).strip()


def normalise_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", normalise(clean(value)))


def normalise_url(value: str) -> str:
    text = clean(value).casefold().split("?", 1)[0].rstrip("/")
    text = text.replace("https://za.linkedin.com/", "https://www.linkedin.com/")
    text = text.replace("http://za.linkedin.com/", "https://www.linkedin.com/")
    text = text.replace("http://www.linkedin.com/", "https://www.linkedin.com/")
    return text


def split_values(value: Any) -> list[str]:
    text = clean(value)
    if not text:
        return []
    values = [clean(part) for part in re.split(r"[\n;]+", text)]
    return [part for part in values if part]


def unique(values: Iterable[str]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        value = clean(value)
        key = normalise(value)
        if value and key and key not in seen:
            seen.add(key)
            result.append(value)
    return result


def header_positions(headers: list[Any]) -> dict[str, list[int]]:
    positions: dict[str, list[int]] = defaultdict(list)
    for index, header in enumerate(headers):
        key = normalise_header(header)
        if key:
            positions[key].append(index)
    return positions


def first_index(positions: dict[str, list[int]], *keys: str) -> int | None:
    for key in keys:
        found = positions.get(normalise_header(key), [])
        if found:
            return found[0]
    return None


def value_at(row: tuple[Any, ...] | list[Any], index: int | None) -> str:
    if index is None or index >= len(row):
        return ""
    return clean(row[index])


def field_values(row: tuple[Any, ...], headers: list[Any], label: str) -> list[str]:
    target = normalise_header(label)
    return unique(value_at(row, index) for index, header in enumerate(headers) if normalise_header(header) == target)


def make_position(
    *,
    title: str = "",
    company: str = "",
    sector: str = "",
    company_size: str = "",
    website: str = "",
    company_linkedin: str = "",
    company_domain: str = "",
    company_description: str = "",
    company_sub_industry: str = "",
    company_country: str = "",
    company_city: str = "",
    company_revenue: str = "",
    company_year_founded: str = "",
    company_specialties: str = "",
) -> dict[str, str]:
    return {
        "title": clean(title),
        "company": clean(company),
        "sector": clean(sector),
        "companySize": clean(company_size),
        "website": clean(website),
        "companyLinkedin": clean(company_linkedin),
        "companyDomain": clean(company_domain),
        "companyDescription": clean(company_description),
        "companySubIndustry": clean(company_sub_industry),
        "companyCountry": clean(company_country),
        "companyCity": clean(company_city),
        "companyRevenue": clean(company_revenue),
        "companyYearFounded": clean(company_year_founded),
        "companySpecialties": clean(company_specialties),
    }


def valid_name(name: str) -> bool:
    key = normalise(name)
    return bool(key and key not in {"first name", "last name", "name", "manager from 1 november 2024"})


def phone_entries(number: Any, phone_type: str) -> list[dict[str, str]]:
    entries = []
    for value in split_values(number):
        if len(re.sub(r"\D", "", value)) >= 7:
            entries.append({"number": value, "type": phone_type})
    return entries


def email_entries(address: Any, status: str, email_type: str) -> list[dict[str, Any]]:
    entries = []
    for value in split_values(address):
        if "@" not in value:
            continue
        entries.append({
            "address": value,
            "status": clean(status),
            "type": email_type,
            "masked": value.startswith("...@") or "***" in value,
        })
    return entries


def base_record(name: str, file_name: str, sheet: str, row_number: int) -> dict[str, Any]:
    return {
        "name": clean(name),
        "positions": [],
        "linkedinUrls": [],
        "emails": [],
        "phones": [],
        "locations": [],
        "seniorities": [],
        "departments": [],
        "technologies": [],
        "software": [],
        "tags": [],
        "triggers": [],
        "remarks": [],
        "outreachStatuses": [],
        "datesAdded": [],
        "sourceSheets": [sheet],
        "sourceRecords": [{"file": file_name, "sheet": sheet, "row": row_number}],
        "relatedRoles": [],
    }


def extract_workbook(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    records: list[dict[str, Any]] = []
    ignored_sheets: list[dict[str, Any]] = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    file_name = path.name

    for worksheet in workbook.worksheets:
        rows = worksheet.iter_rows(values_only=True)
        first_row = tuple(next(rows, ()))
        sheet = clean(worksheet.title)
        headers = list(first_row)
        positions = header_positions(headers)

        if sheet == "AECI" and "firstname" not in positions:
            all_rows = [first_row, *rows]
            for row_number, row in enumerate(all_rows, start=1):
                name = f"{value_at(row, 0)} {value_at(row, 1)}".strip()
                if not valid_name(name):
                    continue
                record = base_record(name, file_name, sheet, row_number)
                record["positions"].append(make_position(
                    title=value_at(row, 2), company=value_at(row, 3) or sheet,
                ))
                records.append(record)
            continue

        if sheet == "Africa Tech Fest":
            for row_number, row in enumerate(rows, start=2):
                name = value_at(row, 0)
                if not valid_name(name):
                    continue
                record = base_record(name, file_name, sheet, row_number)
                record["positions"].append(make_position(title=value_at(row, 1), company=value_at(row, 2)))
                linkedin = normalise_url(value_at(row, 3))
                if linkedin:
                    record["linkedinUrls"].append(linkedin)
                record["emails"].extend(email_entries(value_at(row, 4), "", "work"))
                record["phones"].extend(phone_entries(value_at(row, 5), "mobile"))
                records.append(record)
            continue

        if sheet == "BMW":
            secondary_positions: dict[str, list[int]] | None = None
            secondary_company = ""
            for row_number, row in enumerate(rows, start=2):
                if normalise_header(value_at(row, 0)) == "firstname" and normalise_header(value_at(row, 1)) == "lastname":
                    secondary_positions = header_positions(list(row))
                    continue

                if secondary_positions is not None:
                    first_name = value_at(row, first_index(secondary_positions, "First Name"))
                    last_name = value_at(row, first_index(secondary_positions, "Last Name"))
                    name = f"{first_name} {last_name}".strip()
                    if not valid_name(name):
                        continue
                    secondary_company = value_at(row, first_index(secondary_positions, "Company Name")) or secondary_company
                    record = base_record(name, file_name, sheet, row_number)
                    record["positions"].append(make_position(
                        title=value_at(row, first_index(secondary_positions, "Job Title")),
                        company=secondary_company,
                    ))
                    linkedin = normalise_url(value_at(row, first_index(secondary_positions, "LinkedIn URL")))
                    if linkedin:
                        record["linkedinUrls"].append(linkedin)
                    email_status = value_at(row, first_index(secondary_positions, "Email Status (ZB)"))
                    record["emails"].extend(email_entries(
                        value_at(row, first_index(secondary_positions, "Email Address")),
                        email_status,
                        "work",
                    ))
                    for label, kind in (("Business Phone Number", "business"), ("Direct Phone Number", "direct"), ("Mobile Number", "mobile")):
                        record["phones"].extend(phone_entries(value_at(row, first_index(secondary_positions, label)), kind))
                    remark = value_at(row, first_index(secondary_positions, "Remarks"))
                    if remark:
                        record["remarks"].append(remark)
                    date_added = value_at(row, first_index(secondary_positions, "Date Added"))
                    if date_added:
                        record["datesAdded"].append(date_added)
                    records.append(record)
                    continue

                name = value_at(row, first_index(positions, "Manager from 1 November 2024"))
                if not valid_name(name):
                    continue
                record = base_record(name, file_name, sheet, row_number)
                record["positions"].append(make_position(company="BMW"))
                linkedin = normalise_url(value_at(row, first_index(positions, "LinkedIn URL")))
                if linkedin:
                    record["linkedinUrls"].append(linkedin)
                email_status = value_at(row, first_index(positions, "Email Status (ZB)"))
                record["emails"].extend(email_entries(value_at(row, first_index(positions, "Email Address")), email_status, "work"))
                for label, kind in (("Business Phone Number", "business"), ("Direct Phone Number", "direct"), ("Mobile Number", "mobile")):
                    record["phones"].extend(phone_entries(value_at(row, first_index(positions, label)), kind))
                record["remarks"] = unique(field_values(row, headers, "Notes") + field_values(row, headers, "Remarks"))
                record["datesAdded"] = field_values(row, headers, "Date Added")
                related = {
                    "roleTitle": value_at(row, first_index(positions, "Role Title")),
                    "roleName": value_at(row, first_index(positions, "Role Name")),
                    "roleGroup": value_at(row, first_index(positions, "Role Group")),
                    "level": value_at(row, first_index(positions, "Level")),
                    "newRoleGroup": value_at(row, first_index(positions, "New Role Group")),
                    "newRoleName": value_at(row, first_index(positions, "New Role Name")),
                    "status": value_at(row, first_index(positions, "Status")),
                    "departmentCode": value_at(row, first_index(positions, "Dept Code from 1 November 2024")),
                    "location": "",
                }
                if any(related.values()):
                    record["relatedRoles"].append(related)
                records.append(record)
            continue

        first_name_index = first_index(positions, "First Name")
        last_name_index = first_index(positions, "Last Name")
        if last_name_index is not None and first_name_index is None and last_name_index == 1:
            first_name_index = 0

        full_name_index = first_index(positions, "Name")
        if first_name_index is None and last_name_index is None and full_name_index is None:
            ignored_sheets.append({"file": file_name, "sheet": sheet, "reason": "no person-name columns"})
            continue

        title_index = first_index(positions, "Job Title")
        company_index = first_index(positions, "Company Name")
        website_index = first_index(positions, "Website")
        company_linkedin_index = first_index(positions, "Company LI Profile")
        sector_index = first_index(positions, "Industry")
        size_index = first_index(positions, "Employee Size", "Company Size")
        location_index = first_index(positions, "Location") if sheet == "Java AWS SA" else None
        software_index = first_index(positions, "Software")
        context = {"company": "", "website": "", "companyLinkedin": "", "sector": "", "companySize": ""}

        for row_number, row in enumerate(rows, start=2):
            if company_index is not None:
                context["company"] = value_at(row, company_index) or context["company"]
            elif sheet not in {"PWC", "The Coca-Cola Company"}:
                context["company"] = sheet
            context["website"] = value_at(row, website_index) or context["website"]
            context["companyLinkedin"] = value_at(row, company_linkedin_index) or context["companyLinkedin"]
            context["sector"] = value_at(row, sector_index) or context["sector"]
            context["companySize"] = value_at(row, size_index) or context["companySize"]

            if full_name_index is not None and first_name_index is None:
                name = value_at(row, full_name_index)
            else:
                name = f"{value_at(row, first_name_index)} {value_at(row, last_name_index)}".strip()
            if not valid_name(name):
                continue

            company = value_at(row, company_index) or context["company"]
            record = base_record(name, file_name, sheet, row_number)
            record["positions"].append(make_position(
                title=value_at(row, title_index),
                company=company,
                sector=context["sector"],
                company_size=context["companySize"],
                website=context["website"],
                company_linkedin=context["companyLinkedin"],
            ))

            linkedin = normalise_url(value_at(row, first_index(positions, "LinkedIn URL")))
            if linkedin:
                record["linkedinUrls"].append(linkedin)
            email_status = value_at(row, first_index(positions, "Email Status (ZB)"))
            record["emails"].extend(email_entries(value_at(row, first_index(positions, "Email Address")), email_status, "work"))
            for label, kind in (("Business Phone Number", "business"), ("Direct Phone Number", "direct"), ("Mobile Number", "mobile")):
                record["phones"].extend(phone_entries(value_at(row, first_index(positions, label)), kind))
            if location_index is not None:
                location = value_at(row, location_index)
                if location:
                    record["locations"].append(location)
            software = value_at(row, software_index)
            if software:
                record["software"].append(software)
            trigger = value_at(row, first_index(positions, "Trigger"))
            if trigger:
                record["triggers"].append(trigger)
            record["remarks"] = unique(field_values(row, headers, "Remarks") + field_values(row, headers, "Notes"))
            record["outreachStatuses"] = unique(field_values(row, headers, "Status") + field_values(row, headers, "Interested?"))
            record["datesAdded"] = field_values(row, headers, "Date Added")

            job_posting = value_at(row, first_index(positions, "Job Posting"))
            job_location = value_at(row, first_index(positions, "Location")) if sheet == "M-Kopa" else ""
            if job_posting or job_location:
                record["relatedRoles"].append({
                    "roleTitle": job_posting,
                    "roleName": "",
                    "roleGroup": "",
                    "level": "",
                    "newRoleGroup": "",
                    "newRoleName": "",
                    "status": "",
                    "departmentCode": "",
                    "location": job_location,
                })
            records.append(record)

    return records, ignored_sheets


def extract_csv(path: Path) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    with path.open(newline="", encoding="utf-8-sig") as handle:
        for row_number, row in enumerate(csv.DictReader(handle), start=2):
            name = f"{clean(row.get('First Name'))} {clean(row.get('Last Name'))}".strip()
            if not valid_name(name):
                continue
            record = base_record(name, path.name, "CSV export", row_number)
            location = unique([row.get("City", ""), row.get("State", ""), row.get("Country", "")])
            if location:
                record["locations"].append(", ".join(location))
            linkedin = normalise_url(clean(row.get("LinkedIn URL")))
            if linkedin:
                record["linkedinUrls"].append(linkedin)
            for address_key, status_key, email_type in (
                ("Work Email", "Work Email Confidence", "work"),
                ("Direct Email", "Direct Email Confidence", "direct"),
                ("Additional Email 1", "Additional Email 1 Confidence", "additional"),
                ("Additional Email 2", "Additional Email 2 Confidence", "additional"),
            ):
                record["emails"].extend(email_entries(row.get(address_key), clean(row.get(status_key)), email_type))
            for number_key, type_key in (("Phone 1", "Phone 1 Type"), ("Phone 2", "Phone 2 Type")):
                record["phones"].extend(phone_entries(row.get(number_key), clean(row.get(type_key)) or "phone"))
            record["positions"].append(make_position(
                title=clean(row.get("Job Title")),
                company=clean(row.get("Company Name")),
                sector=clean(row.get("Company Main Industry")),
                company_size=clean(row.get("Company Number of Employees")),
                website=clean(row.get("Company Website")),
                company_linkedin=clean(row.get("Company linkedin URL")),
                company_domain=clean(row.get("Company Domain")),
                company_description=clean(row.get("Company Description")),
                company_sub_industry=clean(row.get("Company Sub Industry")),
                company_country=clean(row.get("Company Country")),
                company_city=clean(row.get("Company City")),
                company_revenue=clean(row.get("Company Revenue")),
                company_year_founded=clean(row.get("Company Year Founded")),
                company_specialties=clean(row.get("Company Specialties")),
            ))
            record["seniorities"] = unique([row.get("Seniority", "")])
            record["departments"] = unique([row.get("Departments", "")])
            record["technologies"] = unique(split_values(row.get("Company Technologies")))
            record["tags"] = unique(split_values(row.get("Tags")))
            record["datesAdded"] = unique([row.get("Date", "")])
            records.append(record)
    return records


@dataclass
class DisjointSet:
    parent: list[int]

    @classmethod
    def create(cls, size: int) -> "DisjointSet":
        return cls(list(range(size)))

    def find(self, item: int) -> int:
        while self.parent[item] != item:
            self.parent[item] = self.parent[self.parent[item]]
            item = self.parent[item]
        return item

    def union(self, left: int, right: int) -> None:
        left_root = self.find(left)
        right_root = self.find(right)
        if left_root != right_root:
            self.parent[right_root] = left_root


def strong_email(record: dict[str, Any]) -> list[str]:
    return [entry["address"].casefold() for entry in record["emails"] if not entry["masked"]]


def group_records(records: list[dict[str, Any]]) -> list[list[dict[str, Any]]]:
    dsu = DisjointSet.create(len(records))
    linkedin_owner: dict[str, int] = {}
    name_company_title_owner: dict[str, int] = {}
    email_name_owner: dict[str, int] = {}
    mobile_name_owner: dict[str, int] = {}

    for index, record in enumerate(records):
        name_key = normalise(record["name"])
        positions = record["positions"]
        for linkedin in record["linkedinUrls"]:
            key = normalise_url(linkedin)
            if "/in/" not in key:
                continue
            if key in linkedin_owner:
                dsu.union(index, linkedin_owner[key])
            else:
                linkedin_owner[key] = index
        for position in positions:
            company_key = normalise(position["company"])
            title_key = normalise(position["title"])
            key = f"{name_key}|{company_key}|{title_key}"
            if name_key and company_key:
                if key in name_company_title_owner:
                    dsu.union(index, name_company_title_owner[key])
                else:
                    name_company_title_owner[key] = index
        for email in strong_email(record):
            key = f"{name_key}|{email}"
            if key in email_name_owner:
                dsu.union(index, email_name_owner[key])
            else:
                email_name_owner[key] = index
        for phone in record["phones"]:
            if phone["type"].casefold() != "mobile":
                continue
            digits = re.sub(r"\D", "", phone["number"])
            if len(digits) < 9:
                continue
            key = f"{name_key}|{digits}"
            if key in mobile_name_owner:
                dsu.union(index, mobile_name_owner[key])
            else:
                mobile_name_owner[key] = index

    grouped: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for index, record in enumerate(records):
        grouped[dsu.find(index)].append(record)
    return list(grouped.values())


def unique_dicts(values: Iterable[dict[str, Any]], key_fields: tuple[str, ...]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    seen: set[tuple[str, ...]] = set()
    for value in values:
        key = tuple(normalise(str(value.get(field, ""))) for field in key_fields)
        if key not in seen:
            seen.add(key)
            result.append(value)
    return result


def merge_group(group: list[dict[str, Any]]) -> dict[str, Any]:
    names = unique(record["name"] for record in group)
    name = sorted(names, key=lambda value: (-len(value), value.casefold()))[0]
    merged: dict[str, Any] = {
        "id": "",
        "name": name,
        "nameAliases": [value for value in names if normalise(value) != normalise(name)],
    }
    list_fields = (
        "linkedinUrls", "locations", "seniorities", "departments", "technologies", "software",
        "tags", "triggers", "remarks", "outreachStatuses", "datesAdded", "sourceSheets",
    )
    for field in list_fields:
        merged[field] = unique(value for record in group for value in record[field])
    merged["positions"] = unique_dicts(
        (value for record in group for value in record["positions"]),
        ("title", "company", "sector", "companySize", "website", "companyLinkedin"),
    )
    merged["emails"] = unique_dicts(
        (value for record in group for value in record["emails"]),
        ("address", "status", "type"),
    )
    merged["phones"] = unique_dicts(
        (value for record in group for value in record["phones"]),
        ("number", "type"),
    )
    merged["sourceRecords"] = unique_dicts(
        (value for record in group for value in record["sourceRecords"]),
        ("file", "sheet", "row"),
    )
    merged["relatedRoles"] = unique_dicts(
        (value for record in group for value in record["relatedRoles"]),
        ("roleTitle", "roleName", "roleGroup", "level", "newRoleGroup", "newRoleName", "status", "departmentCode", "location"),
    )
    return merged


def assign_ids(contacts: list[dict[str, Any]]) -> None:
    slug_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for contact in contacts:
        slug = re.sub(r"[^a-z0-9]+", "-", normalise(contact["name"])).strip("-") or "contact"
        slug_groups[slug].append(contact)
    for slug, group in slug_groups.items():
        if len(group) == 1:
            group[0]["id"] = slug
            continue
        for contact in group:
            identity = "|".join(contact["linkedinUrls"] + [email["address"] for email in contact["emails"]] + [position["company"] for position in contact["positions"]])
            suffix = hashlib.sha1(identity.encode()).hexdigest()[:8]
            contact["id"] = f"{slug}-{suffix}"


def enrich_exact_company_metadata(contacts: list[dict[str, Any]]) -> None:
    metadata: dict[str, dict[str, str]] = defaultdict(dict)
    fields = (
        "sector", "companySize", "website", "companyLinkedin", "companyDomain", "companyDescription",
        "companySubIndustry", "companyCountry", "companyCity", "companyRevenue", "companyYearFounded",
        "companySpecialties",
    )
    for contact in contacts:
        for position in contact["positions"]:
            company_key = normalise(position["company"])
            if not company_key:
                continue
            for field in fields:
                if position[field] and field not in metadata[company_key]:
                    metadata[company_key][field] = position[field]
    for contact in contacts:
        for position in contact["positions"]:
            company_key = normalise(position["company"])
            for field, value in metadata.get(company_key, {}).items():
                if not position[field]:
                    position[field] = value


def validate(raw_records: list[dict[str, Any]], contacts: list[dict[str, Any]]) -> dict[str, Any]:
    source_record_count = sum(len(contact["sourceRecords"]) for contact in contacts)
    duplicate_ids = len(contacts) - len({contact["id"] for contact in contacts})
    contacts_without_company = sum(not any(position["company"] for position in contact["positions"]) for contact in contacts)
    contacts_without_title = sum(not any(position["title"] for position in contact["positions"]) for contact in contacts)
    contacts_without_location = sum(not contact["locations"] for contact in contacts)
    return {
        "rawPersonRows": len(raw_records),
        "uniqueContacts": len(contacts),
        "mergedDuplicateRows": len(raw_records) - len(contacts),
        "sourceRecordCount": source_record_count,
        "sourceRowsReconciled": source_record_count == len(raw_records),
        "duplicateIds": duplicate_ids,
        "contactsWithoutCompany": contacts_without_company,
        "contactsWithoutTitle": contacts_without_title,
        "contactsWithoutLocation": contacts_without_location,
        "companies": len({normalise(position["company"]) for contact in contacts for position in contact["positions"] if position["company"]}),
        "titles": len({normalise(position["title"]) for contact in contacts for position in contact["positions"] if position["title"]}),
        "sectors": len({normalise(position["sector"]) for contact in contacts for position in contact["positions"] if position["sector"]}),
        "locations": len({normalise(location) for contact in contacts for location in contact["locations"] if location}),
    }


def main() -> None:
    if len(sys.argv) < 4:
        raise SystemExit("Usage: build_contact_directory.py OUTPUT_DIRECTORY AUDIT_JSON SOURCE...")
    output_path = Path(sys.argv[1])
    audit_path = Path(sys.argv[2])
    sources = [Path(value) for value in sys.argv[3:]]
    raw_records: list[dict[str, Any]] = []
    ignored_sheets: list[dict[str, Any]] = []
    source_counts: dict[str, int] = {}

    for source in sources:
        if source.suffix.casefold() == ".csv":
            records = extract_csv(source)
            ignored = []
        else:
            records, ignored = extract_workbook(source)
        raw_records.extend(records)
        ignored_sheets.extend(ignored)
        source_counts[source.name] = len(records)

    contacts = [merge_group(group) for group in group_records(raw_records)]
    enrich_exact_company_metadata(contacts)
    assign_ids(contacts)
    contacts.sort(key=lambda contact: (normalise(contact["name"]), contact["id"]))
    audit = validate(raw_records, contacts)
    audit["sourcePersonRows"] = source_counts
    audit["ignoredSheets"] = ignored_sheets

    if not audit["sourceRowsReconciled"] or audit["duplicateIds"]:
        raise RuntimeError(f"Contact reconciliation failed: {audit}")

    output_path.mkdir(parents=True, exist_ok=True)
    for stale_part in output_path.glob("contacts-*.json"):
        stale_part.unlink()
    chunk_size = 400
    part_names: list[str] = []
    for index, offset in enumerate(range(0, len(contacts), chunk_size), start=1):
        part_name = f"contacts-{index:03d}.json"
        part_names.append(part_name)
        (output_path / part_name).write_text(
            json.dumps(contacts[offset:offset + chunk_size], ensure_ascii=False, separators=(",", ":")),
            encoding="utf-8",
        )
    imports = [f"import part{index:03d} from './{part_name}'" for index, part_name in enumerate(part_names, start=1)]
    spreads = ", ".join(f"...part{index:03d}" for index in range(1, len(part_names) + 1))
    (output_path / "index.ts").write_text(
        "\n".join([*imports, "", f"export const contactsRaw = [{spreads}]", ""]),
        encoding="utf-8",
    )
    audit_path.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
