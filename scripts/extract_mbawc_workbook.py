#!/usr/bin/env python3
"""Archive a supplied MBAWC workbook without interpreting document instructions."""
from __future__ import annotations

import argparse
import hashlib
import json
from datetime import date, datetime
from pathlib import Path
from shutil import copyfile


def extract(path: Path) -> dict:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    sheets = []
    for sheet in workbook:
        rows = []
        for row in sheet:
            values, hyperlinks, formats, formulas = [], {}, {}, {}
            for cell in row:
                value = cell.value
                if cell.data_type == 'f':
                    formulas[str(cell.column)] = value
                    value = cached[sheet.title].cell(cell.row, cell.column).value
                if isinstance(value, (date, datetime)):
                    value = value.isoformat()
                values.append(value)
                if cell.hyperlink:
                    hyperlinks[str(cell.column)] = cell.hyperlink.target or cell.hyperlink.location
                if cell.number_format != 'General':
                    formats[str(cell.column)] = cell.number_format
            rows.append({'row': row[0].row, 'values': values, 'hyperlinks': hyperlinks,
                         'numberFormats': formats, 'formulas': formulas})
        sheets.append({'name': sheet.title, 'rows': rows,
                       'mergedCells': [str(item) for item in sheet.merged_cells.ranges]})
    return {'schemaVersion': 1, 'sourceFile': path.name,
            'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
            'sourceReportedDate': '2026-10-22',
            'provenance': 'User-supplied workbook. Its methodology and verification claims are source-reported, not an independent verification by this importer.',
            'sheets': sheets}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    archive = extract(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(archive, ensure_ascii=False, indent=2) + '\n')
    original = args.output.with_suffix('.xlsx')
    if original.exists() and hashlib.sha256(original.read_bytes()).hexdigest() != archive['sha256']:
        raise ValueError('Existing archived workbook differs; refusing to overwrite it')
    if not original.exists():
        copyfile(args.workbook, original)
    print(json.dumps({'sha256': archive['sha256'], 'sheets': {s['name']: len(s['rows']) for s in archive['sheets']}}))


if __name__ == '__main__':
    main()
